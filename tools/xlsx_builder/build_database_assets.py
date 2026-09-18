#!/usr/bin/env python3
"""Build and verify RusMorph runtime assets from read-only XLSX/CSV/DOCX sources."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    from tools.xlsx_builder import GENERATOR_VERSION, NORMALIZATION_VERSION, SCHEMA_VERSION
    from tools.xlsx_builder.declension_parser import parse_declension_sheets
    from tools.xlsx_builder.knowledge_association import KnowledgeAssociationEngine, association_report
    from tools.xlsx_builder.lexicon_parser import parse_lexicon
    from tools.xlsx_builder.models import BuildIssue, KnowledgeChunk, LexiconRecord
    from tools.xlsx_builder.russian_normalizer import normalize_russian_for_search
    from tools.xlsx_builder.validation import ValidationError, verify_assets
    from tools.xlsx_builder.workbook_auditor import audit_sources, classify_sheet, discover_sources, source_sha256
    from tools.xlsx_builder.openrussian_importer import build_openrussian_enhanced_lexicon
else:
    from . import GENERATOR_VERSION, NORMALIZATION_VERSION, SCHEMA_VERSION
    from .declension_parser import parse_declension_sheets
    from .knowledge_association import KnowledgeAssociationEngine, association_report
    from .lexicon_parser import parse_lexicon
    from .models import BuildIssue, KnowledgeChunk, LexiconRecord
    from .russian_normalizer import normalize_russian_for_search
    from .validation import ValidationError, verify_assets
    from .workbook_auditor import audit_sources, classify_sheet, discover_sources, source_sha256
    from .openrussian_importer import build_openrussian_enhanced_lexicon

RUSSIAN_TOKEN = re.compile(r"[А-Яа-яЁё]+(?:\u0301[А-Яа-яЁё]*)*")
EXAMPLE_SPLIT = re.compile(r"(?:\r?\n)+|[•●]")
ALTERNATION_PATTERN = re.compile(r"[А-Яа-яЁё]{1,3}\s*[-–—]\s*[А-Яа-яЁё]{1,3}")
CATEGORY_BY_PREFIX = {
    "1.": "OTHER",
    "2.": "OTHER",
    "3.": "STEM_ALTERNATION",
    "4.": "STEM_ALTERNATION",
    "5.": "OTHER",
}


def _stable_id(prefix: str, *parts: str) -> str:
    digest = hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()[:20]
    return f"{prefix}_{digest}"


def parse_knowledge_documents(paths: list[Path]) -> tuple[list[KnowledgeChunk], list[BuildIssue]]:
    chunks: list[KnowledgeChunk] = []
    issues: list[BuildIssue] = []
    for path in paths:
        try:
            document = Document(path)
        except (OSError, ValueError) as exc:
            issues.append(BuildIssue("ERROR", "DOCX_READ_FAILED", str(exc), path.name))
            continue
        for table_index, table in enumerate(document.tables, 1):
            if not table.rows:
                continue
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            for row_index, row in enumerate(table.rows[1:], 2):
                values = [cell.text.strip() for cell in row.cells]
                if not any(values):
                    continue
                title = values[0] if values else f"表 {table_index} 第 {row_index} 行"
                description = values[1] if len(values) > 1 else ""
                history = values[2] if len(values) > 2 else ""
                examples_text = values[3] if len(values) > 3 else ""
                examples = [item.strip() for item in EXAMPLE_SPLIT.split(examples_text) if item.strip()]
                content_parts = [part for part in (description, history) if part]
                category = next((value for prefix, value in CATEGORY_BY_PREFIX.items() if title.startswith(prefix)), "GENERAL_PROJECT_BACKGROUND")
                russian_keywords = list(dict.fromkeys(
                    normalize_russian_for_search(token)
                    for token in RUSSIAN_TOKEN.findall(examples_text)
                    if len(normalize_russian_for_search(token)) >= 2
                ))
                keywords = list(dict.fromkeys([title.split("(", 1)[0].strip(), *russian_keywords]))
                lemma_patterns = list(dict.fromkeys(
                    normalize_russian_for_search(match.group(0))
                    for example in examples
                    if (match := RUSSIAN_TOKEN.search(example)) is not None
                    and not example[max(0, match.start() - 1):match.start()].endswith("*")
                ))
                alternations = list(dict.fromkeys(
                    re.sub(r"\s+", "", match.group(0)).replace("–", "-").replace("—", "-")
                    for match in ALTERNATION_PATTERN.finditer(" ".join(content_parts + examples))
                ))
                morphology_tags = [title.split("(", 1)[0].lstrip("0123456789. ").strip()]
                chunks.append(KnowledgeChunk(
                    id=_stable_id("knowledge", category, title, "\n".join(content_parts)),
                    title=title, category=category, keywords=keywords,
                    content="\n\n".join(content_parts), examples=examples,
                    alternation_patterns=alternations,
                    lemma_patterns=lemma_patterns,
                    morphology_tags=[item for item in morphology_tags if item],
                    source_document=path.name,
                    section_path=["词法解释", headers[0] if headers else f"表 {table_index}", title],
                ))
        if not document.tables:
            issues.append(BuildIssue("WARNING", "DOCX_NO_TABLES", "DOCX 中没有可解析表格", path.name))
    return chunks, issues


def _source_version(paths: list[Path], configuration_paths: list[Path]) -> str:
    digest = hashlib.sha256()
    digest.update(f"schema={SCHEMA_VERSION};generator={GENERATOR_VERSION};normalization={NORMALIZATION_VERSION}".encode())
    for path in sorted([*paths, *configuration_paths], key=lambda item: (item.parent.name.casefold(), item.name.casefold())):
        digest.update(path.name.encode("utf-8"))
        digest.update(_configuration_sha(path).encode("ascii") if path in configuration_paths else source_sha256(path).encode("ascii"))
    return digest.hexdigest()


def _configuration_sha(path: Path) -> str:
    if path.is_file():
        return source_sha256(path)
    empty_overrides = b'{"entryMappings":[],"lemmaMappings":[]}'
    return hashlib.sha256(empty_overrides).hexdigest()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _read_json(path: Path, default: Any) -> Any:
    if not path.is_file():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValidationError(f"无法解析配置 {path}: {exc}") from exc


def _agent_record(entry: LexiconRecord) -> dict[str, Any]:
    morphology = {
        key: value for key, value in {
            "gender": entry.gender, "declensionClass": entry.declension_class,
            "endingType": entry.ending_type, "pluralStressPattern": entry.plural_stress_pattern,
            "aspect": entry.aspect, "conjugationClass": entry.conjugation_class,
            "phoneticAlternation": entry.phonetic_alternation,
        }.items() if value is not None
    }
    fragments = [f"俄语词：{entry.display_form}"]
    if entry.chinese_meaning:
        fragments.append(f"中文：{entry.chinese_meaning}")
    if entry.parts_of_speech:
        fragments.append(f"词性：{'、'.join(entry.parts_of_speech)}")
    labels = {
        "gender": "性", "declensionClass": "变格法", "endingType": "结尾类型",
        "pluralStressPattern": "复数重音变化", "aspect": "体",
        "conjugationClass": "变位法", "phoneticAlternation": "语音交替",
    }
    fragments.extend(f"{labels[key]}：{value}" for key, value in morphology.items())
    if entry.annotations:
        fragments.extend(
            f"{item.field_name}：{item.value}"
            for item in entry.annotations
            if item.field_name not in {"词类", "词类2", "词类3", "词类4", *labels.values()}
        )
    source_fields = [key for key in morphology]
    knowledge = [
        {
            "chunkId": item.knowledge_chunk_id,
            "category": item.category,
            "title": item.title,
            "score": item.score,
            "associationType": item.association_type,
            "reason": item.reason,
        }
        for item in entry.knowledge_associations
    ]
    source = entry.provenance[0]
    retrieval_text = "词表字段：" + "；".join(fragments) + "。"
    if knowledge:
        retrieval_text += "相关知识类别：" + "、".join(dict.fromkeys(item["category"] for item in knowledge)) + "。"
    return {
        "id": entry.id, "headword": entry.display_form,
        "normalizedHeadword": entry.normalized_primary_form,
        "searchForms": list(dict.fromkeys(item.normalized for item in entry.search_forms)),
        **({"meaningZh": entry.chinese_meaning} if entry.chinese_meaning else {}),
        **({"partsOfSpeech": entry.parts_of_speech} if entry.parts_of_speech else {}),
        **({"rawPartsOfSpeech": entry.raw_parts_of_speech} if entry.raw_parts_of_speech != entry.parts_of_speech else {}),
        **({"morphology": morphology} if morphology else {}),
        "knowledge": knowledge,
        "localEvidence": {
            "hasStructuredMorphology": bool(source_fields),
            "hasKnowledgeExplanation": bool(knowledge),
            "sourceFields": source_fields,
        },
        "retrievalText": retrieval_text,
        "source": {"file": source.source_file, "sheet": source.source_sheet, "row": source.source_row},
    }


def _entry_fixture(entry: LexiconRecord | None) -> dict[str, Any]:
    if entry is None:
        return {"available": False, "entry": None}
    return {
        "available": True,
        "entry": {
            "id": entry.id, "displayForm": entry.display_form,
            "normalizedLemma": entry.normalized_primary_form,
            "searchForms": [item.normalized for item in entry.search_forms],
            "partsOfSpeech": entry.parts_of_speech,
            "rawPartsOfSpeech": entry.raw_parts_of_speech,
            "sourceRow": entry.provenance[0].source_row,
        },
    }


def real_asset_fixtures(entries: list[LexiconRecord]) -> dict[str, Any]:
    duplicate_groups: dict[str, list[LexiconRecord]] = {}
    for entry in entries:
        duplicate_groups.setdefault(entry.normalized_primary_form, []).append(entry)
    duplicate = next((items for _, items in sorted(duplicate_groups.items()) if len(items) > 1), [])
    return {
        "totalEntries": len(entries),
        "fieldNonEmptyCounts": {
            "partsOfSpeech": sum(bool(item.parts_of_speech) for item in entries),
            "gender": sum(bool(item.gender) for item in entries),
            "declensionClass": sum(bool(item.declension_class) for item in entries),
            "endingType": sum(bool(item.ending_type) for item in entries),
            "pluralStressPattern": sum(bool(item.plural_stress_pattern) for item in entries),
            "aspect": sum(bool(item.aspect) for item in entries),
            "conjugationClass": sum(bool(item.conjugation_class) for item in entries),
            "phoneticAlternation": sum(bool(item.phonetic_alternation) for item in entries),
        },
        "samples": {
            "phoneticAlternation": _entry_fixture(next((item for item in entries if item.phonetic_alternation), None)),
            "conjugationClass": _entry_fixture(next((item for item in entries if item.conjugation_class), None)),
            "aspect": _entry_fixture(next((item for item in entries if item.aspect), None)),
            "multipleSearchForms": _entry_fixture(next((item for item in entries if len(item.search_forms) > 1), None)),
            "multiplePartsOfSpeech": _entry_fixture(next((item for item in entries if len(item.parts_of_speech) > 1), None)),
            "accentedDisplayForm": _entry_fixture(next((item for item in entries if "\u0301" in item.display_form), None)),
            "legalDuplicateLemma": {
                "available": bool(duplicate),
                "entries": [_entry_fixture(item)["entry"] for item in duplicate[:4]],
            },
            "connectorAliases": [
                _entry_fixture(item)["entry"]
                for item in entries if "连接词" in item.raw_parts_of_speech
            ],
        },
    }


def build_assets(input_dir: Path, output_dir: Path, agent_output: Path, report_dir: Path) -> dict[str, Any]:
    audit = audit_sources(input_dir)
    config_dir = Path(__file__).resolve().parent / "config"
    alias_path = config_dir / "part_of_speech_aliases.json"
    rules_path = config_dir / "knowledge_association_rules.json"
    prefixes_path = config_dir / "russian_prefixes.json"
    overrides_path = input_dir / "knowledge_overrides.json"
    aliases = _read_json(alias_path, {})
    rules = _read_json(rules_path, {})
    prefixes = _read_json(prefixes_path, {"prefixes": []})
    overrides = _read_json(overrides_path, {"entryMappings": [], "lemmaMappings": []})
    lexicon = parse_lexicon(audit.sheets, aliases)
    declension = parse_declension_sheets(audit.sheets)
    _, _, docx_paths = discover_sources(input_dir)
    knowledge, knowledge_issues = parse_knowledge_documents(docx_paths)
    all_issues = [*audit.issues, *lexicon.issues, *declension.issues, *knowledge_issues]
    morphology_coverage = real_asset_fixtures(lexicon.entries)["fieldNonEmptyCounts"]
    if lexicon.entries and morphology_coverage["phoneticAlternation"] == 0:
        all_issues.append(BuildIssue(
            "WARNING",
            "PHONETIC_ALTERNATION_NOT_ANNOTATED",
            "主词表的语音交替字段没有已标注词条；涉及音变的条件检索将准确返回空集，不使用模型补造。",
        ))
    if any(issue.level == "ERROR" for issue in all_issues):
        audit.report.update({"warnings": [item.to_dict() for item in all_issues if item.level == "WARNING"], "errors": [item.to_dict() for item in all_issues if item.level == "ERROR"]})
        _write_json(report_dir / "xlsx_audit_report.json", audit.report)
        raise ValidationError("源数据包含阻止构建的 ERROR，请查看 xlsx_audit_report.json")

    association = KnowledgeAssociationEngine(rules, prefixes.get("prefixes", []), overrides).associate(lexicon.entries, knowledge)
    selected_sheet_keys = {
        (item["file"], item["sheet"]) for item in lexicon.selected_sheets
    }
    source_paths = sorted(
        {
            sheet.source
            for sheet in audit.sheets
            if (sheet.source.name, sheet.name) in selected_sheet_keys
            or classify_sheet(sheet) == "DECLENSION"
        } | set(docx_paths),
        key=lambda path: path.name.casefold(),
    )
    configuration_paths = [alias_path, rules_path, prefixes_path, overrides_path]
    version = _source_version(source_paths, configuration_paths)
    lexicon_assets = [item.to_asset() for item in lexicon.entries]
    declension_asset = {
        "schemaVersion": SCHEMA_VERSION, "generatorVersion": GENERATOR_VERSION,
        "normalizationVersion": NORMALIZATION_VERSION, "dataVersion": version,
        "rules": [item.to_asset() for item in declension.rules],
        "rawData": [], "warnings": declension.warnings,
    }
    knowledge_assets = [item.to_asset() for item in knowledge]
    counts = {
        "entries": len(lexicon_assets),
        "searchForms": sum(len(item["searchForms"]) for item in lexicon_assets),
        "partsOfSpeechRelations": sum(len(item["partsOfSpeech"]) for item in lexicon_assets),
        "annotationRecords": sum(len(item["annotations"]) for item in lexicon_assets),
        "provenanceRecords": sum(len(item["provenance"]) for item in lexicon_assets),
        "declensionRules": len(declension.rules), "knowledgeChunks": len(knowledge_assets),
        "crossRefs": len(association.associations),
    }
    manifest = {
        "schemaVersion": SCHEMA_VERSION,
        "sourceFiles": [{"name": path.name, "sha256": source_sha256(path), "size": path.stat().st_size} for path in source_paths],
        "configurationFiles": [
            {"name": path.name, "sha256": _configuration_sha(path), "size": path.stat().st_size if path.is_file() else 0}
            for path in configuration_paths
        ],
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "generatorVersion": GENERATOR_VERSION,
        "normalizationVersion": NORMALIZATION_VERSION,
        "dataVersion": version, "counts": counts,
        "morphologyCoverage": morphology_coverage,
        "warnings": sum(issue.level == "WARNING" for issue in all_issues) + len(declension.warnings),
    }
    _write_json(output_dir / "lexicon.json", lexicon_assets)
    _write_json(output_dir / "declension_rules.json", declension_asset)
    _write_json(output_dir / "knowledge_chunks.json", knowledge_assets)
    _write_json(output_dir / "data_manifest.json", manifest)
    _write_json(report_dir / "pure_word_validation.json", lexicon.pure_word_validation)
    _write_json(report_dir / "knowledge_association_report.json", association_report(lexicon.entries, knowledge, association))
    _write_json(report_dir / "manual_review_candidates.json", association.manual_review_candidates)
    _write_json(report_dir / "real_asset_test_fixtures.json", real_asset_fixtures(lexicon.entries))
    audit.report.update({
        "detectedMainSheets": lexicon.selected_sheets,
        "selectedMainSheets": lexicon.selected_sheets,
        "skippedDuplicateMainSheets": lexicon.skipped_duplicate_sheets,
        "entryCount": counts["entries"], "multiFormCellCount": lexicon.multi_form_cells,
        "legalDuplicateWordCount": lexicon.legal_duplicate_words,
        "exactDuplicateRecordCount": lexicon.duplicate_records,
        "rawPartOfSpeechValues": lexicon.raw_part_of_speech_values,
        "normalizedPartOfSpeechValues": lexicon.normalized_part_of_speech_values,
        "appliedAliases": lexicon.applied_aliases,
        "unknownPartOfSpeechValues": lexicon.unknown_part_of_speech_values,
        "unknownPartsOfSpeech": sorted({issue.message for issue in all_issues if issue.code == "UNKNOWN_PART_OF_SPEECH"}),
        "fieldNonEmptyCounts": morphology_coverage,
        "unrecognizedDeclensionRules": declension.warnings,
        "pureWordDifferences": lexicon.pure_word_validation,
        "warnings": [item.to_dict() for item in all_issues if item.level == "WARNING"],
        "errors": [item.to_dict() for item in all_issues if item.level == "ERROR"],
        "info": [item.to_dict() for item in all_issues if item.level == "INFO"],
    })
    _write_json(report_dir / "xlsx_audit_report.json", audit.report)
    agent_output.mkdir(parents=True, exist_ok=True)
    with (agent_output / "lexicon_agent.jsonl").open("w", encoding="utf-8", newline="\n") as handle:
        for entry in lexicon.entries:
            handle.write(json.dumps(_agent_record(entry), ensure_ascii=False, separators=(",", ":")) + "\n")
    openrussian_dir = input_dir / "openrussian"
    if openrussian_dir.is_dir() and any(openrussian_dir.glob("*.csv")):
        build_openrussian_enhanced_lexicon(output_dir, openrussian_dir)
        manifest = json.loads((output_dir / "data_manifest.json").read_text(encoding="utf-8"))
        counts = manifest.get("counts", counts)
    verify_assets(output_dir)
    return {"counts": counts, "manifest": manifest, "audit": audit.report}


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=root / "data-source")
    parser.add_argument("--output", type=Path, default=root / "app/src/main/assets/database")
    parser.add_argument("--agent-output", type=Path, default=root / "build/generated/agent")
    parser.add_argument("--report-output", type=Path, default=root / "build/reports/data")
    parser.add_argument("--verify-only", action="store_true")
    parser.add_argument("--allow-empty", action="store_true")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.verify_only:
            result = verify_assets(args.output, allow_empty=args.allow_empty)
            print(json.dumps(result, ensure_ascii=False))
            return 0
        xlsx, csv_files, _ = discover_sources(args.input)
        if not xlsx and not csv_files:
            if all((args.output / name).is_file() for name in ("lexicon.json", "declension_rules.json", "knowledge_chunks.json", "data_manifest.json")):
                print("WARNING: 未找到 XLSX/CSV，保留并验证现有 assets。", file=sys.stderr)
                verify_assets(args.output, allow_empty=args.allow_empty)
                return 0
            raise ValidationError("data-source 中没有 XLSX/CSV，且不存在完整的预生成 assets")
        result = build_assets(args.input, args.output, args.agent_output, args.report_output)
        print(json.dumps(result["counts"], ensure_ascii=False))
        return 0
    except (OSError, ValidationError, ValueError) as exc:
        print(f"数据构建失败：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
