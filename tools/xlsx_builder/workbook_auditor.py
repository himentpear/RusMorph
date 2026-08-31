"""Read-only XLSX/CSV discovery, classification, and audit reporting."""

from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import BuildIssue

REQUIRED_LEXICON_HEADERS = {"课号", "序号", "俄语", "中文", "词类"}
CASE_LABELS = {"主格", "属格", "与格", "宾格", "工具格", "前置格"}
IGNORED_HEADER_RE = re.compile(r"^(?:column\d+|unnamed(?::\s*\d+)?)$", re.IGNORECASE)


@dataclass
class SheetData:
    source: Path
    name: str
    rows: list[list[Any]]
    original_max_row: int
    original_max_column: int

    @property
    def effective_rows(self) -> list[list[Any]]:
        return [row for _, row in self.effective_indexed_rows]

    @property
    def effective_indexed_rows(self) -> list[tuple[int, list[Any]]]:
        return [
            (index, row)
            for index, row in enumerate(self.rows, 1)
            if any(_has_value(value) for value in row)
        ]

    @property
    def effective_max_row(self) -> int:
        indices = [index for index, row in enumerate(self.rows, 1) if any(_has_value(v) for v in row)]
        return max(indices, default=0)


@dataclass
class WorkbookAudit:
    sheets: list[SheetData]
    report: dict[str, Any]
    issues: list[BuildIssue]


def _has_value(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def normalized_header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def discover_sources(input_dir: Path) -> tuple[list[Path], list[Path], list[Path]]:
    xlsx = sorted(input_dir.glob("*.xlsx"), key=lambda item: item.name.casefold())
    csv_files = sorted(input_dir.glob("*.csv"), key=lambda item: item.name.casefold())
    docx = sorted(input_dir.glob("*.docx"), key=lambda item: item.name.casefold())
    return (
        [path for path in xlsx if not path.name.startswith("~$")],
        csv_files,
        [path for path in docx if not path.name.startswith("~$")],
    )


def _read_xlsx(path: Path) -> list[SheetData]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise ValueError(f"无法读取 {path.name}: {exc}") from exc
    result: list[SheetData] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            result.append(
                SheetData(path, worksheet.title, rows, worksheet.max_row, worksheet.max_column)
            )
    finally:
        workbook.close()
    return result


def _read_csv(path: Path) -> SheetData:
    data = path.read_bytes()
    text: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError(f"无法识别 {path.name} 的文本编码")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(text.splitlines(), dialect=dialect))
    maximum = max((len(row) for row in rows), default=0)
    return SheetData(path, path.stem, rows, len(rows), maximum)


def load_sources(input_dir: Path) -> tuple[list[SheetData], list[Path]]:
    xlsx, csv_files, docx = discover_sources(input_dir)
    sheets = [sheet for path in xlsx for sheet in _read_xlsx(path)]
    sheets.extend(_read_csv(path) for path in csv_files)
    return sheets, docx


def classify_sheet(sheet: SheetData) -> str:
    effective = sheet.effective_rows
    if not effective:
        return "EMPTY"
    header = {normalized_header(value) for value in effective[0] if _has_value(value)}
    flattened = {normalized_header(value) for row in effective[:25] for value in row if _has_value(value)}
    if len(REQUIRED_LEXICON_HEADERS & header) >= 4 and "俄语" in header:
        return "LEXICON"
    if len(CASE_LABELS & flattened) >= 4:
        return "DECLENSION"
    first = effective[0]
    if len(first) >= 3 and normalized_header(first[0]) not in {"课号", "lesson"}:
        numeric_pairs = sum(
            _looks_integer(row[0]) and len(row) > 2 and _looks_integer(row[1])
            for row in effective[:30]
        )
        required_pairs = 1 if len(effective) == 1 else max(2, min(10, len(effective)) // 2)
        if numeric_pairs >= required_pairs:
            return "PURE_WORDS"
    return "OTHER"


def _looks_integer(value: Any) -> bool:
    try:
        return float(str(value).strip()).is_integer()
    except (TypeError, ValueError):
        return False


def canonical_rows(sheet: SheetData) -> set[str]:
    if classify_sheet(sheet) != "LEXICON" or not sheet.effective_rows:
        return set()
    header = [normalized_header(value) for value in sheet.effective_rows[0]]
    useful_indices = [
        index for index, value in enumerate(header)
        if value and not IGNORED_HEADER_RE.match(value)
    ]
    result: set[str] = set()
    for row in sheet.effective_rows[1:]:
        values = [str(row[index]).strip() if index < len(row) and row[index] is not None else "" for index in useful_indices]
        if any(values):
            result.add("\x1f".join(values))
    return result


def find_duplicate_candidates(sheets: list[SheetData]) -> list[dict[str, Any]]:
    candidates = [sheet for sheet in sheets if classify_sheet(sheet) == "LEXICON"]
    duplicates: list[dict[str, Any]] = []
    for left_index, left in enumerate(candidates):
        left_rows = canonical_rows(left)
        for right in candidates[left_index + 1 :]:
            right_rows = canonical_rows(right)
            denominator = min(len(left_rows), len(right_rows))
            overlap = len(left_rows & right_rows) / denominator if denominator else 0.0
            if overlap >= 0.90:
                duplicates.append({
                    "left": {"file": left.source.name, "sheet": left.name},
                    "right": {"file": right.source.name, "sheet": right.name},
                    "overlapRatio": round(overlap, 6),
                })
    return duplicates


def source_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def audit_sources(input_dir: Path) -> WorkbookAudit:
    issues: list[BuildIssue] = []
    try:
        sheets, _ = load_sources(input_dir)
    except (OSError, ValueError) as exc:
        issue = BuildIssue("ERROR", "SOURCE_READ_FAILED", str(exc))
        return WorkbookAudit([], {"files": [], "sheets": [], "errors": [issue.to_dict()]}, [issue])

    sheet_reports: list[dict[str, Any]] = []
    for sheet in sheets:
        effective = sheet.effective_rows
        classification = classify_sheet(sheet)
        headers = effective[0] if effective and classification == "LEXICON" else []
        ignored_columns = [
            index + 1 for index, header in enumerate(headers)
            if not normalized_header(header) or IGNORED_HEADER_RE.match(normalized_header(header))
        ]
        sheet_reports.append({
            "file": sheet.source.name,
            "name": sheet.name,
            "classification": classification,
            "originalMaxRow": sheet.original_max_row,
            "effectiveMaxRow": sheet.effective_max_row,
            "validDataRows": max(0, len(effective) - (1 if classification == "LEXICON" else 0)),
            "originalMaxColumn": sheet.original_max_column,
            "headers": headers,
            "blankTrailingRows": max(0, sheet.original_max_row - sheet.effective_max_row),
            "ignoredColumns": ignored_columns,
        })
    duplicate_sheets = find_duplicate_candidates(sheets)
    for item in duplicate_sheets:
        issues.append(BuildIssue("INFO", "DUPLICATE_SHEET", "检测到高度重合的主词表", str(item)))
    files = sorted({sheet.source for sheet in sheets}, key=lambda path: path.name.casefold())
    _, _, docx = discover_sources(input_dir)
    files.extend(path for path in docx if path not in files)
    report = {
        "files": [{"name": path.name, "size": path.stat().st_size, "sha256": source_sha256(path)} for path in files],
        "sheets": sheet_reports,
        "duplicateSheets": duplicate_sheets,
        "warnings": [item.to_dict() for item in issues if item.level == "WARNING"],
        "errors": [item.to_dict() for item in issues if item.level == "ERROR"],
        "info": [item.to_dict() for item in issues if item.level == "INFO"],
    }
    return WorkbookAudit(sheets, report, issues)
