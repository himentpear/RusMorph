from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from tools.data_builder.build_data import (
    build_all,
    clean_value,
    normalize_russian_for_search,
    russian_word_matches,
)


WORKBOOK_XML = """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
 <sheets>
  <sheet name="大学俄语（一）" sheetId="1" r:id="rId1"/>
  <sheet name="大学俄语（一） (2)" sheetId="2" r:id="rId2"/>
  <sheet name="变格表" sheetId="3" r:id="rId3"/>
 </sheets>
</workbook>
"""

RELATIONSHIPS_XML = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
 <Relationship Id="rId1" Target="worksheets/sheet1.xml"
  Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/>
 <Relationship Id="rId2" Target="worksheets/sheet2.xml"
  Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/>
 <Relationship Id="rId3" Target="worksheets/sheet3.xml"
  Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/>
</Relationships>
"""


def inline_cell(reference: str, value: str) -> str:
    return f'<c r="{reference}" t="inlineStr"><is><t>{value}</t></is></c>'


MAIN_SHEET_XML = f"""<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
 <sheetData>
  <row r="1">{inline_cell("A1", "课号")}{inline_cell("B1", "序号")}
   {inline_cell("C1", "俄语")}{inline_cell("D1", "中文")}
   {inline_cell("E1", "词类")}{inline_cell("F1", "语音交替")}</row>
  <row r="2">{inline_cell("A2", "18")}{inline_cell("B2", "18")}
   {inline_cell("C2", "автомоби́ль")}{inline_cell("D2", "汽车")}
   {inline_cell("E2", "名词")}{inline_cell("F2", "8")}</row>
  <row r="3">{inline_cell("A3", "18")}{inline_cell("B3", "19")}
   {inline_cell("C3", "писа́ть, пишу́, пи́шешь, пи́шут")}{inline_cell("D3", "写")}
   {inline_cell("E3", "动词")}{inline_cell("F3", "с-ш")}</row>
 </sheetData>
</worksheet>
"""

DUPLICATE_SHEET_XML = f"""<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
 <sheetData>
  <row r="1">{inline_cell("A1", "课号")}{inline_cell("B1", "俄语")}</row>
  <row r="2">{inline_cell("A2", "18")}{inline_cell("B2", "автомоби́ль")}</row>
 </sheetData>
</worksheet>
"""

DECLENSION_SHEET_XML = f"""<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
 <sheetData>
  <row r="1">{inline_cell("A1", "名词 单数 阳性 主格")}{inline_cell("B1", "-а")}</row>
  <row r="2">{inline_cell("A2", "无法解释的注释")}</row>
 </sheetData>
</worksheet>
"""


class NormalizationTest(unittest.TestCase):
    def test_required_normalization_and_boundaries(self) -> None:
        self.assertEqual(normalize_russian_for_search("писа́ть"), "писать")
        self.assertTrue(russian_word_matches("автомобиль", "автомоби́ль"))
        self.assertTrue(russian_word_matches("елка", "ёлка"))
        self.assertTrue(russian_word_matches("ёлка", "елка"))
        self.assertTrue(
            russian_word_matches("писать", "писа́ть, пишу́, пи́шешь, пи́шут")
        )
        self.assertFalse(russian_word_matches("есть", "шесть"))

    def test_string_and_numeric_eight_are_null(self) -> None:
        self.assertIsNone(clean_value("8"))
        self.assertIsNone(clean_value(8))
        self.assertIsNone(clean_value(8.0))


class IntegrationBuildTest(unittest.TestCase):
    def test_builds_assets_and_links_iotation(self) -> None:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover
            self.skipTest(f"python-docx unavailable: {exc}")

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "data-source"
            assets = root / "assets"
            report_path = root / "report.json"
            source.mkdir()

            with zipfile.ZipFile(source / "lexicon.xlsx", "w") as archive:
                archive.writestr("xl/workbook.xml", WORKBOOK_XML)
                archive.writestr("xl/_rels/workbook.xml.rels", RELATIONSHIPS_XML)
                archive.writestr("xl/worksheets/sheet1.xml", MAIN_SHEET_XML)
                archive.writestr("xl/worksheets/sheet2.xml", DUPLICATE_SHEET_XML)
                archive.writestr("xl/worksheets/sheet3.xml", DECLENSION_SHEET_XML)

            document = Document()
            document.add_heading("动词变位中的音变现象", level=1)
            document.add_heading("j-音组变化", level=2)
            document.add_paragraph("с-ш 与 б-бл：писать → пишу，любить → люблю。")
            document.save(source / "knowledge.docx")

            report = build_all(source, assets, report_path)
            first_asset_bytes = {
                path.name: path.read_bytes() for path in sorted(assets.iterdir())
            }
            second_report = build_all(source, assets, report_path)
            second_asset_bytes = {
                path.name: path.read_bytes() for path in sorted(assets.iterdir())
            }
            lexicon = json.loads((assets / "lexicon.json").read_text(encoding="utf-8"))
            rules = json.loads(
                (assets / "declension_rules.json").read_text(encoding="utf-8")
            )
            chunks = json.loads(
                (assets / "knowledge_chunks.json").read_text(encoding="utf-8")
            )

            self.assertEqual(report["totalLexiconEntries"], 2)
            self.assertEqual(report, second_report)
            self.assertEqual(first_asset_bytes, second_asset_bytes)
            automobile = next(item for item in lexicon if item["lemma"] == "автомобиль")
            writing = next(item for item in lexicon if item["lemma"] == "писать")
            self.assertIn("автомобиль", automobile["searchForms"])
            self.assertEqual(writing["searchForms"][:2], ["писать", "пишу"])
            self.assertTrue(writing["relatedKnowledgeChunkIds"])
            self.assertEqual(chunks[0]["category"], "IOTATION")
            self.assertEqual(len(rules["rules"]), 1)
            self.assertTrue(rules["rawData"])


if __name__ == "__main__":
    unittest.main()
