#!/usr/bin/env python3
"""Build RusMorph Android JSON assets from read-only XLSX and DOCX sources."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

try:
    from tools.data_builder.audit_sources import (
        AuditError,
        SheetData,
        column_letter,
        read_sheet,
        shared_strings,
        sheet_fingerprint,
        source_files,
        worksheet_targets,
    )
except ModuleNotFoundError:  # Allow direct execution from tools/data_builder.
    from audit_sources import (  # type: ignore[no-redef]
        AuditError,
        SheetData,
        column_letter,
        read_sheet,
        shared_strings,
        sheet_fingerprint,
        source_files,
        worksheet_targets,
    )


ACCENT = "\u0301"
FORM_SEPARATOR = re.compile(r"\s*[,，;；]\s*")
SPACE = re.compile(r"\s+")
DUPLICATE_SUFFIX = re.compile(r"\s*\(\d+\)\s*$")
RUSSIAN_WORD = re.compile(r"[А-Яа-яЁё]+(?:\u0301[А-Яа-яЁё]*)*")
ALTERNATION = re.compile(r"[А-Яа-яЁё]{1,2}\s*[-–—]\s*[А-Яа-яЁё]{1,3}")
EXAMPLE_ARROW = re.compile(
    r"[А-Яа-яЁё\u0301-]+(?:\s+[А-Яа-яЁё\u0301-]+){0,3}"
    r"\s*(?:→|->|⇒)\s*"
    r"[А-Яа-яЁё\u0301-]+(?:\s+[А-Яа-яЁё\u0301-]+){0,3}"
)


HEADER_ALIASES: dict[str, str] = {
    "课号": "lesson",
    "序号": "sequence",
    "俄语": "displayForm",
    "中文": "chineseMeaning",
    "词类": "partOfSpeech1",
    "词类1": "partOfSpeech1",
    "词类2": "partOfSpeech2",
    "词类3": "partOfSpeech3",
    "词类4": "partOfSpeech4",
    "性": "gender",
    "变格法": "declensionClass",
    "结尾字母": "endingType",
    "复数重音转移": "pluralStressPattern",
    "动词的体": "aspect",
    "变位法": "conjugationClass",
    "语音交替": "phoneticAlternation",
}

LEXICON_FIELDS = (
    "lesson",
    "sequence",
    "displayForm",
    "chineseMeaning",
    "partOfSpeech1",
    "partOfSpeech2",
    "partOfSpeech3",
    "partOfSpeech4",
    "gender",
    "declensionClass",
    "endingType",
    "pluralStressPattern",
    "aspect",
    "conjugationClass",
    "phoneticAlternation",
)

CATEGORY_PATTERNS: list[tuple[str, tuple[str, ...]]] = [
    ("ATHEMATIC_CONJUGATION", ("无主题", "非主题", "атемат", "дать", "есть")),
    ("FIRST_PALATALIZATION", ("第一腭化", "г-ж", "к-ч", "х-ш")),
    ("IOTATION", ("j-音", "j音", "йотац", "iotation", "с-ш", "д-ж", "б-бл")),
    ("MIXED_CONJUGATION", ("混合变位", "特殊变位", "хотеть")),
    ("VA_SUFFIX_LOSS", ("ва 后缀", "ва后缀", "-ва-", "-ва")),
    ("NASAL_VOWEL_REMAINS", ("鼻化元音", "鼻音元音", "nasal vowel")),
    ("PAST_TENSE_HISTORY", ("过去时历史", "过去时的历史", "历史过去时")),
    ("FUTURE_TENSE_HISTORY", ("将来时历史", "将来时的历史", "历史将来时")),
    ("STEM_ALTERNATION", ("词干交替", "音位交替", "чередован", "交替现象")),
    ("GENERAL_PROJECT_BACKGROUND", ("项目背景", "工程背景", "总体背景", "概述")),
]

CATEGORY_KEYWORDS: dict[str, list[str]] = {
    "ATHEMATIC_CONJUGATION": ["无主题变位", "дать", "есть"],
    "FIRST_PALATALIZATION": ["第一腭化", "г-ж", "к-ч", "х-ш"],
    "IOTATION": ["j音组", "с-ш", "д-ж", "б-бл"],
    "MIXED_CONJUGATION": ["混合变位", "хотеть"],
    "VA_SUFFIX_LOSS": ["ва后缀", "-ва-"],
    "NASAL_VOWEL_REMAINS": ["鼻化元音"],
    "STEM_ALTERNATION": ["词干交替"],
    "PAST_TENSE_HISTORY": ["过去时", "历史"],
    "FUTURE_TENSE_HISTORY": ["将来时", "历史"],
    "GENERAL_PROJECT_BACKGROUND": ["项目背景"],
}

RULE_CATEGORIES: list[tuple[str, tuple[str, ...]]] = [
    ("NOUN_DECLENSION", ("名词",)),
    ("ADJECTIVE_HARD", ("形容词硬音", "硬音形容词", "形容词", "硬音")),
    ("ADJECTIVE_SOFT", ("形容词软音", "软音形容词", "形容词", "软音")),
    ("SEVEN_LETTER_RULE", ("七字母", "7字母")),
    ("FIVE_LETTER_RULE", ("五字母", "5字母")),
    ("TSE_SPELLING_RULE", ("ц", "拼写规则")),
]

CASE_NAMES = ("主格", "属格", "与格", "宾格", "工具格", "前置格", "介词格")
NUMBER_LABELS = (("复数", "PLURAL"), ("单数", "SINGULAR"))
GENDER_LABELS = (("阳性", "MASCULINE"), ("阴性", "FEMININE"), ("中性", "NEUTER"))


@dataclass
class BuildStats:
    duplicate_entries: int = 0
    discarded_records: int = 0
    unknown_fields: int = 0
    source_warnings: list[str] = field(default_factory=list)


def clean_value(value: Any) -> Any | None:
    """Convert all specified empty/8 sentinels to None and trim strings."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value == 8:
        return None
    if isinstance(value, str):
        cleaned = unicodedata.normalize("NFC", value).strip()
        if not cleaned or cleaned == "8":
            return None
        return cleaned
    return value


def normalize_russian_for_search(text: str) -> str:
    """Normalize Russian text for search without damaging ь, ъ, й, or ё."""
    normalized = unicodedata.normalize("NFC", text).lower()
    normalized = normalized.replace(ACCENT, "")
    return SPACE.sub(" ", normalized).strip()


def russian_search_variants(text: str) -> list[str]:
    """Return stable ё-preserving and е-folded variants for one form."""
    base = normalize_russian_for_search(text)
    if not base:
        return []
    variants = [base]
    folded = base.replace("ё", "е")
    if folded != base:
        variants.append(folded)
    return variants


def split_russian_forms(text: str) -> list[str]:
    """Extract comma/semicolon-separated forms while preserving display text."""
    forms: list[str] = []
    for raw_form in FORM_SEPARATOR.split(unicodedata.normalize("NFC", text)):
        form = raw_form.strip()
        if form and any(character.isalpha() for character in form):
            forms.append(form)
    return forms


def russian_word_matches(query: str, candidate: str) -> bool:
    """Match normalized Russian using Unicode word boundaries, never substrings."""
    candidate_variants = russian_search_variants(candidate)
    for query_variant in russian_search_variants(query):
        folded_query = query_variant.replace("ё", "е")
        pattern = re.compile(rf"(?<!\w){re.escape(folded_query)}(?!\w)", re.UNICODE)
        for candidate_variant in candidate_variants:
            if pattern.search(candidate_variant.replace("ё", "е")):
                return True
    return False


def stable_id(prefix: str, *parts: Any) -> str:
    """Create a deterministic ID from normalized, order-stable business keys."""
    canonical = "\x1f".join(
        normalize_russian_for_search(str(part)) if part is not None else ""
        for part in parts
    )
    digest = hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:20]
    return f"{prefix}_{digest}"


def normalized_header(value: str) -> str:
    """Normalize a source header conservatively for alias matching."""
    return SPACE.sub("", unicodedata.normalize("NFC", value)).strip()


def parse_integer(value: Any) -> int | str | None:
    """Return integral source values as int and preserve non-integral labels."""
    cleaned = clean_value(value)
    if cleaned is None:
        return None
    text = str(cleaned).strip()
    try:
        numeric = float(text)
        if numeric.is_integer():
            return int(numeric)
    except ValueError:
        pass
    return text


def unique_nonempty(values: Iterable[Any]) -> list[str]:
    """Deduplicate cleaned strings while preserving their first occurrence."""
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = clean_value(value)
        if cleaned is None:
            continue
        text = str(cleaned)
        key = normalize_russian_for_search(text)
        if key not in seen:
            seen.add(key)
            result.append(text)
    return result


def choose_header_row(sheet: SheetData) -> tuple[int | None, dict[int, str]]:
    """Choose the row with the most recognized lexicon headers near the top."""
    best_row: int | None = None
    best_mapping: dict[int, str] = {}
    for row_index in sorted(sheet.rows)[:30]:
        mapping: dict[int, str] = {}
        for column, value in sheet.rows[row_index].items():
            field_name = HEADER_ALIASES.get(normalized_header(value))
            if field_name:
                mapping[column] = field_name
        if len(mapping) > len(best_mapping):
            best_row, best_mapping = row_index, mapping
    return best_row, best_mapping


def sheet_completeness(sheet: SheetData, header_row: int, mapping: dict[int, str]) -> int:
    """Score a candidate sheet by cleaned recognized cells below its header."""
    return sum(
        clean_value(row.get(column)) is not None
        for row_index, row in sheet.rows.items()
        if row_index > header_row
        for column in mapping
    )


def load_workbook(path: Path) -> list[SheetData]:
    """Read all worksheets from an XLSX without mutating the archive."""
    try:
        with zipfile.ZipFile(path) as archive:
            strings = shared_strings(archive, path)
            return [
                read_sheet(archive, path, name, target, strings)
                for name, target in worksheet_targets(archive, path)
            ]
    except (OSError, zipfile.BadZipFile, AuditError) as exc:
        raise AuditError(f"Could not read {path.name}: {exc}") from exc


def select_lexicon_sheets(
    workbook: Path,
    sheets: Sequence[SheetData],
    stats: BuildStats,
) -> list[tuple[SheetData, int, dict[int, str]]]:
    """Keep the most complete candidate from each near-duplicate sheet group."""
    candidates: list[tuple[SheetData, int, dict[int, str], int]] = []
    for sheet in sheets:
        if "变格表" in sheet.name:
            continue
        header_row, mapping = choose_header_row(sheet)
        if header_row is None or "displayForm" not in mapping.values():
            continue
        score = sheet_completeness(sheet, header_row, mapping)
        candidates.append((sheet, header_row, mapping, score))

        header_values = sheet.rows.get(header_row, {})
        unknown = [
            value
            for value in header_values.values()
            if clean_value(value) is not None
            and normalized_header(value) not in HEADER_ALIASES
        ]
        stats.unknown_fields += len(unknown)
        if unknown:
            stats.source_warnings.append(
                f"{workbook.name}/{sheet.name}: unrecognized headers {unknown}"
            )

    grouped: defaultdict[str, list[tuple[SheetData, int, dict[int, str], int]]] = defaultdict(list)
    for candidate in candidates:
        base = DUPLICATE_SUFFIX.sub("", normalized_header(candidate[0].name)).casefold()
        grouped[base].append(candidate)

    selected: list[tuple[SheetData, int, dict[int, str]]] = []
    seen_fingerprints: set[str] = set()
    for group in grouped.values():
        group.sort(key=lambda item: (-item[3], item[0].name.casefold()))
        winner = group[0]
        if len(group) > 1:
            skipped = [item[0].name for item in group[1:]]
            stats.source_warnings.append(
                f"{workbook.name}: kept more complete sheet {winner[0].name!r}; "
                f"skipped near-duplicates {skipped}"
            )
        fingerprint = sheet_fingerprint(winner[0])
        if fingerprint in seen_fingerprints:
            stats.source_warnings.append(
                f"{workbook.name}/{winner[0].name}: skipped exact duplicate sheet"
            )
            continue
        seen_fingerprints.add(fingerprint)
        selected.append((winner[0], winner[1], winner[2]))
    return selected


def source_provenance(workbook: Path, sheet: SheetData, row: int) -> dict[str, Any]:
    """Create a portable provenance record using only source-relative names."""
    return {
        "sourceWorkbook": workbook.name,
        "sourceSheet": sheet.name,
        "sourceRow": row,
    }


def row_to_lexicon_entry(
    workbook: Path,
    sheet: SheetData,
    row_index: int,
    row: dict[int, str],
    mapping: dict[int, str],
    stats: BuildStats,
) -> dict[str, Any] | None:
    """Convert one mapped source row to a cleaned LexiconEntry candidate."""
    values: dict[str, Any] = {field: None for field in LEXICON_FIELDS}
    for column, field_name in mapping.items():
        values[field_name] = clean_value(row.get(column))
    if all(value is None for value in values.values()):
        stats.discarded_records += 1
        return None

    display = values["displayForm"]
    if display is None:
        stats.discarded_records += 1
        stats.source_warnings.append(
            f"{workbook.name}/{sheet.name}!{row_index}: skipped row without Russian form"
        )
        return None
    forms = split_russian_forms(str(display))
    if not forms:
        stats.discarded_records += 1
        stats.source_warnings.append(
            f"{workbook.name}/{sheet.name}!{row_index}: Russian cell has no usable form"
        )
        return None
    lemma_display = forms[0]
    lemma = normalize_russian_for_search(lemma_display)
    search_forms = unique_nonempty(
        variant for form in forms for variant in russian_search_variants(form)
    )
    parts_of_speech = unique_nonempty(
        values[f"partOfSpeech{index}"] for index in range(1, 5)
    )
    provenance = source_provenance(workbook, sheet, row_index)
    return {
        "id": "",
        "lesson": parse_integer(values["lesson"]),
        "sequence": parse_integer(values["sequence"]),
        "displayForm": unicodedata.normalize("NFC", str(display)),
        "lemma": lemma,
        "normalizedLemma": lemma,
        "searchForms": search_forms,
        "chineseMeaning": values["chineseMeaning"],
        "partsOfSpeech": parts_of_speech,
        "gender": values["gender"],
        "declensionClass": values["declensionClass"],
        "endingType": values["endingType"],
        "pluralStressPattern": values["pluralStressPattern"],
        "aspect": values["aspect"],
        "conjugationClass": values["conjugationClass"],
        "phoneticAlternation": values["phoneticAlternation"],
        **provenance,
        "provenance": [provenance],
        "relatedKnowledgeChunkIds": [],
    }


def lexicon_merge_key(entry: dict[str, Any]) -> tuple[Any, ...]:
    """Return the user-specified duplicate merge key."""
    return (
        entry["normalizedLemma"],
        SPACE.sub(" ", str(entry.get("chineseMeaning") or "")).strip().casefold(),
        entry.get("lesson"),
        tuple(sorted(normalize_russian_for_search(item) for item in entry["partsOfSpeech"])),
    )


def merge_lexicon_entry(target: dict[str, Any], incoming: dict[str, Any]) -> None:
    """Merge a duplicate, preferring existing non-empty fields and all sources."""
    scalar_fields = (
        "sequence",
        "displayForm",
        "lemma",
        "gender",
        "declensionClass",
        "endingType",
        "pluralStressPattern",
        "aspect",
        "conjugationClass",
        "phoneticAlternation",
    )
    for field_name in scalar_fields:
        if target.get(field_name) is None and incoming.get(field_name) is not None:
            target[field_name] = incoming[field_name]
    target["searchForms"] = unique_nonempty(target["searchForms"] + incoming["searchForms"])
    target["partsOfSpeech"] = unique_nonempty(
        target["partsOfSpeech"] + incoming["partsOfSpeech"]
    )
    known_sources = {
        (item["sourceWorkbook"], item["sourceSheet"], item["sourceRow"])
        for item in target["provenance"]
    }
    for item in incoming["provenance"]:
        key = (item["sourceWorkbook"], item["sourceSheet"], item["sourceRow"])
        if key not in known_sources:
            target["provenance"].append(item)
            known_sources.add(key)


def build_lexicon(
    xlsx_paths: Sequence[Path], stats: BuildStats
) -> tuple[list[dict[str, Any]], list[tuple[Path, SheetData]]]:
    """Build and merge lexicon entries; return declension sheets separately."""
    merged: dict[tuple[Any, ...], dict[str, Any]] = {}
    declension_sheets: list[tuple[Path, SheetData]] = []
    for workbook in xlsx_paths:
        sheets = load_workbook(workbook)
        declension_sheets.extend(
            (workbook, sheet) for sheet in sheets if "变格表" in sheet.name
        )
        for sheet, header_row, mapping in select_lexicon_sheets(workbook, sheets, stats):
            for row_index, row in sorted(sheet.rows.items()):
                if row_index <= header_row:
                    continue
                entry = row_to_lexicon_entry(
                    workbook, sheet, row_index, row, mapping, stats
                )
                if entry is None:
                    continue
                key = lexicon_merge_key(entry)
                if key in merged:
                    stats.duplicate_entries += 1
                    merge_lexicon_entry(merged[key], entry)
                else:
                    merged[key] = entry

    entries = list(merged.values())
    for entry in entries:
        entry["id"] = stable_id("lex", *lexicon_merge_key(entry))
    entries.sort(
        key=lambda item: (
            item.get("lesson") is None,
            str(item.get("lesson") or ""),
            item.get("sequence") is None,
            str(item.get("sequence") or ""),
            item["normalizedLemma"],
            item["id"],
        )
    )
    return entries, declension_sheets


def context_text(sheet: SheetData, row_index: int, column: int) -> str:
    """Collect nearby row/column labels for conservative declension parsing."""
    fragments: list[str] = []
    row = sheet.rows.get(row_index, {})
    for index in range(1, column + 1):
        if index in row:
            fragments.append(row[index])
    for index in sorted(sheet.rows):
        if index > row_index:
            break
        value = sheet.rows[index].get(column)
        if value:
            fragments.append(value)
    return " | ".join(unique_nonempty(fragments))


def classify_rule_category(text: str) -> str | None:
    """Recognize only explicitly supported declension/spelling categories."""
    compact = normalized_header(text).casefold()
    if "形容词" in compact and "硬音" in compact:
        return "ADJECTIVE_HARD"
    if "形容词" in compact and "软音" in compact:
        return "ADJECTIVE_SOFT"
    if ("七字母" in compact or "7字母" in compact):
        return "SEVEN_LETTER_RULE"
    if ("五字母" in compact or "5字母" in compact):
        return "FIVE_LETTER_RULE"
    if "ц" in compact and ("规则" in compact or "拼写" in compact):
        return "TSE_SPELLING_RULE"
    if "名词" in compact:
        return "NOUN_DECLENSION"
    return None


def find_label(text: str, choices: Sequence[tuple[str, str]]) -> str | None:
    """Map the first explicit label in contextual text to a canonical value."""
    for source, canonical in choices:
        if source in text:
            return canonical
    return None


def find_case(text: str) -> str | None:
    """Return a canonical Chinese case label from context."""
    for case_name in CASE_NAMES:
        if case_name in text:
            return "前置格" if case_name == "介词格" else case_name
    return None


def likely_result_ending(value: str) -> bool:
    """Conservatively identify a cell that looks like a Russian ending/result."""
    compact = value.strip()
    if not compact or len(compact) > 40:
        return False
    if any(label in compact for label in CASE_NAMES):
        return False
    return bool(re.search(r"[А-Яа-яЁёьъй]", compact))


def build_declension_rules(
    sheets: Sequence[tuple[Path, SheetData]],
) -> dict[str, Any]:
    """Interpret supported declension cells and preserve every uncertain cell."""
    rules: list[dict[str, Any]] = []
    raw_data: list[dict[str, Any]] = []
    warnings: list[str] = []
    for workbook, sheet in sheets:
        interpreted: set[tuple[int, int]] = set()
        raw_start = len(raw_data)
        for row_index, row in sorted(sheet.rows.items()):
            for column, raw_value in sorted(row.items()):
                value = clean_value(raw_value)
                if value is None:
                    continue
                text = str(value)
                context = context_text(sheet, row_index, column)
                category = classify_rule_category(context)
                case_name = find_case(context)
                if category and case_name and likely_result_ending(text):
                    number = find_label(context, NUMBER_LABELS)
                    gender = find_label(context, GENDER_LABELS)
                    ending_type = None
                    for label in ("硬音", "软音", "七字母", "五字母", "ц"):
                        if label in context:
                            ending_type = label
                            break
                    cell = f"{column_letter(column)}{row_index}"
                    rule = {
                        "id": stable_id(
                            "rule", workbook.name, sheet.name, cell, category, case_name
                        ),
                        "category": category,
                        "gender": gender,
                        "endingType": ending_type,
                        "number": number,
                        "caseName": case_name,
                        "resultEnding": text,
                        "description": context,
                        "sourceWorkbook": workbook.name,
                        "sourceSheet": sheet.name,
                        "sourceCellRange": cell,
                    }
                    rules.append(rule)
                    interpreted.add((row_index, column))

        for row_index, row in sorted(sheet.rows.items()):
            for column, raw_value in sorted(row.items()):
                value = clean_value(raw_value)
                if value is None or (row_index, column) in interpreted:
                    continue
                structural_text = str(value)
                if (
                    classify_rule_category(structural_text)
                    or find_case(structural_text)
                    or find_label(structural_text, NUMBER_LABELS)
                    or find_label(structural_text, GENDER_LABELS)
                ):
                    continue
                cell = f"{column_letter(column)}{row_index}"
                raw_data.append(
                    {
                        "sourceWorkbook": workbook.name,
                        "sourceSheet": sheet.name,
                        "sourceCellRange": cell,
                        "rawData": value,
                    }
                )
        sheet_raw_count = len(raw_data) - raw_start
        if sheet_raw_count:
            warnings.append(
                f"{workbook.name}/{sheet.name}: {sheet_raw_count} cells retained as rawData "
                "because their semantics could not be interpreted reliably."
            )
    rules.sort(key=lambda item: item["id"])
    return {
        "schemaVersion": 1,
        "rules": rules,
        "rawData": raw_data,
        "warnings": warnings,
    }


def is_heading(paragraph: Any) -> tuple[bool, int]:
    """Detect Word heading styles and explicit OOXML outline levels."""
    style_name = (getattr(getattr(paragraph, "style", None), "name", "") or "")
    match = re.search(r"(?:heading|标题)\s*(\d+)", style_name, re.IGNORECASE)
    if match:
        return True, max(1, int(match.group(1)))
    properties = getattr(getattr(paragraph, "_p", None), "pPr", None)
    outline = getattr(properties, "outlineLvl", None) if properties is not None else None
    if outline is not None and getattr(outline, "val", None) is not None:
        return True, int(outline.val) + 1
    return False, 0


def classify_knowledge(title: str, content: str) -> str:
    """Classify a knowledge section using explicit multilingual markers."""
    haystack = normalize_russian_for_search(f"{title}\n{content}")
    for category, markers in CATEGORY_PATTERNS:
        if any(knowledge_marker_matches(marker, haystack) for marker in markers):
            return category
    return "GENERAL_PROJECT_BACKGROUND"


def knowledge_marker_matches(marker: str, haystack: str) -> bool:
    """Use word boundaries for standalone Russian markers such as есть."""
    normalized_marker = normalize_russian_for_search(marker)
    if re.fullmatch(r"[а-яё]+", normalized_marker):
        return russian_word_matches(normalized_marker, haystack)
    return normalized_marker in haystack


def extract_examples(content: str) -> list[str]:
    """Extract explicit arrow examples without inventing unstated transformations."""
    return unique_nonempty(match.group(0) for match in EXAMPLE_ARROW.finditer(content))


def knowledge_keywords(category: str, title: str, content: str) -> list[str]:
    """Combine category terms with source alternations and example forms."""
    alternations = [SPACE.sub("", item) for item in ALTERNATION.findall(content)]
    example_words = [word for example in extract_examples(content) for word in RUSSIAN_WORD.findall(example)]
    title_words = [word for word in re.split(r"[\s、，,：:；;（）()]+", title) if len(word) >= 2]
    return unique_nonempty(
        CATEGORY_KEYWORDS.get(category, []) + alternations + example_words + title_words
    )[:24]


def load_docx_chunks(path: Path, warnings: list[str]) -> list[dict[str, Any]]:
    """Read one DOCX with python-docx and split it by headings/semantic groups."""
    try:
        from docx import Document
        from docx.table import Table
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required when DOCX sources exist; install tools/data_builder/requirements.txt"
        ) from exc
    try:
        document = Document(path)
    except (OSError, ValueError) as exc:
        raise RuntimeError(f"Could not read {path.name}: {exc}") from exc

    chunks: list[dict[str, Any]] = []
    section_path: list[str] = []
    current_title = path.stem
    current_content: list[str] = []

    def flush() -> None:
        nonlocal current_content
        content = "\n".join(item for item in current_content if item).strip()
        if not content:
            current_content = []
            return
        title = current_title or path.stem
        category = classify_knowledge(title, content)
        effective_path = section_path.copy() or [title]
        chunks.append(
            {
                "id": stable_id("knowledge", path.name, *effective_path, content),
                "title": title,
                "category": category,
                "keywords": knowledge_keywords(category, title, content),
                "content": unicodedata.normalize("NFC", content),
                "examples": extract_examples(content),
                "sourceDocument": path.name,
                "sectionPath": effective_path,
            }
        )
        current_content = []

    blocks = document.iter_inner_content() if hasattr(document, "iter_inner_content") else document.paragraphs
    for block in blocks:
        if isinstance(block, Table):
            table_lines = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in block.rows
                if any(cell.text.strip() for cell in row.cells)
            ]
            if table_lines:
                current_content.extend(table_lines)
            continue
        text = unicodedata.normalize("NFC", block.text).strip()
        heading, level = is_heading(block)
        if heading and text:
            flush()
            section_path[level - 1 :] = [text]
            current_title = text
        elif text:
            current_content.append(text)
        elif current_content:
            flush()
    flush()
    if not chunks:
        warnings.append(f"{path.name}: no non-empty semantic knowledge chunks found")
    return chunks


def build_knowledge_chunks(
    docx_paths: Sequence[Path], warnings: list[str]
) -> list[dict[str, Any]]:
    """Build deterministically ordered knowledge chunks from every DOCX."""
    chunks: list[dict[str, Any]] = []
    for path in docx_paths:
        chunks.extend(load_docx_chunks(path, warnings))
    chunks.sort(key=lambda item: (item["sourceDocument"].casefold(), item["sectionPath"], item["id"]))
    return chunks


def is_athematic_lemma(lemma: str) -> bool:
    """Recognize дать/есть families without treating шесть as an есть derivative."""
    normalized = normalize_russian_for_search(lemma)
    if normalized == "есть":
        return True
    eat_prefixes = ("съ", "по", "за", "пере", "про", "на", "до", "объ", "вы", "разъ")
    if normalized.endswith("есть") and normalized[: -len("есть")] in eat_prefixes:
        return True
    give_prefixes = (
        "",
        "от",
        "по",
        "пере",
        "про",
        "при",
        "пре",
        "за",
        "вы",
        "на",
        "до",
        "под",
        "раз",
        "соз",
        "из",
        "с",
        "воз",
        "вз",
        "об",
        "у",
    )
    return normalized.endswith("дать") and normalized[: -len("дать")] in give_prefixes


def expected_categories(entry: dict[str, Any]) -> set[str]:
    """Infer only high-confidence special categories for automatic linking."""
    categories: set[str] = set()
    lemma = entry["normalizedLemma"]
    alternation = normalize_russian_for_search(str(entry.get("phoneticAlternation") or ""))
    conjugation = normalize_russian_for_search(str(entry.get("conjugationClass") or ""))
    if lemma in {"писать", "любить"} or any(pair in alternation for pair in ("с-ш", "д-ж", "б-бл")):
        categories.add("IOTATION")
    if lemma == "мочь" or any(pair in alternation for pair in ("г-ж", "к-ч", "х-ш")):
        categories.add("FIRST_PALATALIZATION")
    if lemma == "хотеть" or any(marker in conjugation for marker in ("混合", "特殊")):
        categories.add("MIXED_CONJUGATION")
    if is_athematic_lemma(lemma):
        categories.add("ATHEMATIC_CONJUGATION")
    return categories


def link_knowledge(
    entries: list[dict[str, Any]], chunks: Sequence[dict[str, Any]]
) -> int:
    """Link entries to chunks by fields, forms, keywords, examples, and categories."""
    unlinked_special = 0
    for entry in entries:
        related: set[str] = set()
        expected = expected_categories(entry)
        field_terms = unique_nonempty(
            [
                entry.get("phoneticAlternation"),
                entry.get("conjugationClass"),
                entry.get("declensionClass"),
                entry.get("pluralStressPattern"),
            ]
        )
        for chunk in chunks:
            searchable = "\n".join(
                [chunk["title"], chunk["content"]]
                + chunk["keywords"]
                + chunk["examples"]
            )
            category_match = chunk["category"] in expected
            field_match = any(
                normalize_russian_for_search(term)
                and normalize_russian_for_search(term)
                in normalize_russian_for_search(searchable)
                for term in field_terms
            )
            form_match = any(
                russian_word_matches(form, searchable) for form in entry["searchForms"]
            )
            if category_match or field_match or form_match:
                related.add(chunk["id"])
        entry["relatedKnowledgeChunkIds"] = sorted(related)
        is_special = bool(expected or field_terms)
        if is_special and not related:
            unlinked_special += 1
    return unlinked_special


def write_json(path: Path, value: Any) -> None:
    """Atomically write stable UTF-8 JSON without ASCII escaping."""
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    temporary.replace(path)


def build_all(source_dir: Path, assets_dir: Path, report_path: Path) -> dict[str, Any]:
    """Build all three assets and return the validation report."""
    stats = BuildStats()
    xlsx_paths = list(source_files(source_dir, ".xlsx"))
    docx_paths = list(source_files(source_dir, ".docx"))
    entries, declension_sheets = build_lexicon(xlsx_paths, stats)
    declension = build_declension_rules(declension_sheets)
    knowledge = build_knowledge_chunks(docx_paths, stats.source_warnings)
    unlinked_special = link_knowledge(entries, knowledge)
    category_counts = Counter(chunk["category"] for chunk in knowledge)
    all_categories = [category for category, _ in CATEGORY_PATTERNS]

    if not xlsx_paths and not docx_paths:
        stats.source_warnings.append("No XLSX or DOCX files were found in data-source.")
    report = {
        "schemaVersion": 1,
        "sourceFileCounts": {"xlsx": len(xlsx_paths), "docx": len(docx_paths)},
        "totalLexiconEntries": len(entries),
        "duplicateLexiconEntriesMerged": stats.duplicate_entries,
        "discardedRecords": stats.discarded_records,
        "unrecognizedFields": stats.unknown_fields,
        "unlinkedSpecialEntries": unlinked_special,
        "declensionRuleCount": len(declension["rules"]),
        "uninterpretedDeclensionCellCount": len(declension["rawData"]),
        "knowledgeChunkCountByCategory": {
            category: category_counts.get(category, 0) for category in all_categories
        },
        "warnings": unique_nonempty(stats.source_warnings + declension["warnings"]),
    }
    write_json(assets_dir / "lexicon.json", entries)
    write_json(assets_dir / "declension_rules.json", declension)
    write_json(assets_dir / "knowledge_chunks.json", knowledge)
    write_json(report_path, report)
    return report


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """Parse portable paths relative to the repository by default."""
    project_root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=project_root / "data-source")
    parser.add_argument(
        "--assets-dir",
        type=Path,
        default=project_root / "app" / "src" / "main" / "assets" / "database",
    )
    parser.add_argument(
        "--report",
        type=Path,
        default=project_root / "data_validation_report.json",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    """CLI entry point with contextual errors and nonzero failure status."""
    args = parse_args(argv)
    try:
        args.source_dir.mkdir(parents=True, exist_ok=True)
        report = build_all(args.source_dir, args.assets_dir, args.report)
    except (OSError, AuditError, RuntimeError) as exc:
        print(f"Data build failed: {exc}", file=sys.stderr)
        return 1
    print(
        f"Built {report['totalLexiconEntries']} lexicon entries, "
        f"{report['declensionRuleCount']} declension rules, and "
        f"{sum(report['knowledgeChunkCountByCategory'].values())} knowledge chunks."
    )
    print(f"Validation report: {args.report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
