from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


GRAMMAR_SHEET = "01_语法点"
QUESTION_SHEET = "02_专四题目"
GRAMMAR_HEADERS = ("point_id", "title_zh", "title_ru", "explanation", "example_ru", "example_zh")
QUESTION_HEADERS = (
    "question_id", "point_id", "exam_year", "stem", "option_a",
    "option_b", "option_c", "option_d", "answer", "analysis",
)
EXPECTED_GRAMMAR_POINTS = 16
EXPECTED_QUESTIONS = 48
EXPECTED_MULTI_LINK_QUESTIONS = 3
EXPECTED_GRAMMAR_WITH_QUESTIONS = 8


class ImportValidationError(RuntimeError):
    pass


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def split_point_ids(value: Any) -> list[str]:
    return [item.strip() for item in re.split(r"[\r\n,;，；]+", _text(value)) if item.strip()]


def _source_question_id(value: Any) -> int:
    if isinstance(value, bool):
        raise ImportValidationError(f"Invalid question_id: {value!r}")
    try:
        number = int(value)
    except (TypeError, ValueError) as error:
        raise ImportValidationError(f"Invalid question_id: {value!r}") from error
    return number


def _exam_period(value: Any) -> tuple[int, str]:
    label = _text(value).replace("~", "–")
    years = re.findall(r"\d{4}", label)
    if not years:
        raise ImportValidationError(f"Invalid exam_year: {value!r}")
    return int(years[0]), label


def _question_id(year_label: str, source_question_id: int) -> str:
    normalized_period = re.sub(r"[^0-9]+", "_", year_label).strip("_")
    return f"TEM4_{normalized_period}_{source_question_id}"


@dataclass(frozen=True)
class ImportResult:
    grammar_points: list[dict[str, Any]]
    questions: list[dict[str, Any]]
    links: list[dict[str, Any]]
    audit: dict[str, Any]


def parse_workbook(path: Path) -> ImportResult:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if GRAMMAR_SHEET not in workbook.sheetnames or QUESTION_SHEET not in workbook.sheetnames:
            raise ImportValidationError(
                f"Workbook must contain {GRAMMAR_SHEET!r} and {QUESTION_SHEET!r}; got {workbook.sheetnames!r}"
            )

        grammar_sheet = workbook[GRAMMAR_SHEET]
        question_sheet = workbook[QUESTION_SHEET]
        grammar_header = tuple(_text(cell) for cell in next(grammar_sheet.iter_rows(min_row=1, max_row=1, max_col=6, values_only=True)))
        question_header = tuple(_text(cell) for cell in next(question_sheet.iter_rows(min_row=1, max_row=1, max_col=10, values_only=True)))
        if grammar_header != GRAMMAR_HEADERS:
            raise ImportValidationError(f"Unexpected grammar headers: {grammar_header!r}")
        if question_header != QUESTION_HEADERS:
            raise ImportValidationError(f"Unexpected question headers: {question_header!r}")

        grammar_points: list[dict[str, Any]] = []
        for sort_order, row in enumerate(
            (row for row in grammar_sheet.iter_rows(min_row=2, max_col=6, values_only=True) if _text(row[0])),
            start=1,
        ):
            point_id, title_zh, title_ru, explanation, example_ru, example_zh = row
            grammar_points.append({
                "pointId": _text(point_id),
                "titleZh": _text(title_zh),
                "titleRu": _text(title_ru),
                "explanation": _text(explanation),
                "exampleRu": _text(example_ru) or None,
                "exampleZh": _text(example_zh) or None,
                "parentPointId": None,
                "category": None,
                "sortOrder": sort_order,
                "contentVersion": 1,
            })

        questions: list[dict[str, Any]] = []
        raw_links: list[tuple[str, str]] = []
        for row in question_sheet.iter_rows(min_row=2, max_col=10, values_only=True):
            if not _text(row[0]):
                continue
            source_id = _source_question_id(row[0])
            year, year_label = _exam_period(row[2])
            question_id = _question_id(year_label, source_id)
            point_ids = split_point_ids(row[1])
            questions.append({
                "questionId": question_id,
                "sourceQuestionId": source_id,
                "sourceType": "TEM4_REAL",
                "examYear": year,
                "examYearLabel": year_label,
                "stem": _text(row[3]),
                "optionA": _text(row[4]),
                "optionB": _text(row[5]),
                "optionC": _text(row[6]),
                "optionD": _text(row[7]),
                "answer": _text(row[8]).upper(),
                "analysis": _text(row[9]) or None,
                "difficulty": None,
                "createdAt": None,
            })
            raw_links.extend((question_id, point_id) for point_id in point_ids)

        grammar_ids = {item["pointId"] for item in grammar_points}
        question_ids = {item["questionId"] for item in questions}
        unknown_point_ids = sorted({point_id for _, point_id in raw_links if point_id not in grammar_ids})
        invalid_answers = sorted(item["questionId"] for item in questions if item["answer"] not in {"A", "B", "C", "D"})
        empty_stems = sorted(item["questionId"] for item in questions if not item["stem"])
        missing_analysis = sorted(item["questionId"] for item in questions if not item["analysis"])
        duplicate_question_ids = sorted(key for key, count in Counter(item["questionId"] for item in questions).items() if count > 1)
        duplicate_point_ids = sorted(key for key, count in Counter(item["pointId"] for item in grammar_points).items() if count > 1)
        links_per_question = Counter(question_id for question_id, _ in raw_links)

        links = [{
            "questionId": question_id,
            "pointId": point_id,
            "role": "UNVERIFIED",
            "weight": 1.0,
            "confidence": 0.6,
            "relationSource": "SOURCE_SPREADSHEET",
            "verified": False,
        } for question_id, point_id in raw_links]

        audit = {
            "grammar_point_count": len(grammar_points),
            "question_count": len(questions),
            "link_count": len(links),
            "multi_link_question_count": sum(count > 1 for count in links_per_question.values()),
            "grammar_points_with_questions": len({point_id for _, point_id in raw_links if point_id in grammar_ids}),
            "grammar_points_without_questions": sorted(grammar_ids - {point_id for _, point_id in raw_links}),
            "unknown_point_ids": unknown_point_ids,
            "duplicate_question_ids": duplicate_question_ids,
            "duplicate_point_ids": duplicate_point_ids,
            "invalid_answers": invalid_answers,
            "empty_stems": empty_stems,
            "missing_analysis": missing_analysis,
            "orphan_question_links": sorted(question_id for question_id, _ in raw_links if question_id not in question_ids),
            "relation_policy": "SOURCE_SPREADSHEET mappings are preserved as UNVERIFIED; no automatic relabeling is applied.",
            "known_relation_quality_notes": [{
                "questionId": "TEM4_2018_8",
                "pointId": "SYN_0008",
                "note": "The source analysis names participles while the source relation points to the relative-clause explanation path. Preserved without automatic correction.",
            }],
        }
        _validate(audit)
        return ImportResult(grammar_points, questions, links, audit)
    finally:
        workbook.close()


def _validate(audit: dict[str, Any]) -> None:
    expected = {
        "grammar_point_count": EXPECTED_GRAMMAR_POINTS,
        "question_count": EXPECTED_QUESTIONS,
        "multi_link_question_count": EXPECTED_MULTI_LINK_QUESTIONS,
        "grammar_points_with_questions": EXPECTED_GRAMMAR_WITH_QUESTIONS,
    }
    mismatches = {key: (audit[key], value) for key, value in expected.items() if audit[key] != value}
    blocking_lists = (
        "unknown_point_ids", "duplicate_question_ids", "duplicate_point_ids",
        "invalid_answers", "empty_stems", "orphan_question_links",
    )
    failures = {key: audit[key] for key in blocking_lists if audit[key]}
    if mismatches or failures:
        raise ImportValidationError(f"Grammar/TEM4 audit failed: mismatches={mismatches}, failures={failures}")


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _update_manifest(manifest_path: Path, source: Path, result: ImportResult, assets: list[Path]) -> None:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source_record = {"name": source.name, "sha256": _sha256(source), "size": source.stat().st_size}
    source_files = [item for item in manifest.get("sourceFiles", []) if item.get("name") != source.name]
    source_files.append(source_record)
    manifest["sourceFiles"] = sorted(source_files, key=lambda item: item["name"].casefold())
    counts = dict(manifest.get("counts", {}))
    counts.update({
        "grammarPoints": len(result.grammar_points),
        "questions": len(result.questions),
        "grammarQuestionLinks": len(result.links),
        "multiLinkQuestions": result.audit["multi_link_question_count"],
    })
    manifest["counts"] = counts
    manifest["grammarTem4AssetHashes"] = {path.name: _sha256(path) for path in assets}
    digest = hashlib.sha256()
    digest.update(str(manifest.get("dataVersion", "")).encode("utf-8"))
    digest.update(source_record["sha256"].encode("ascii"))
    for path in sorted(assets, key=lambda item: item.name):
        digest.update(path.name.encode("utf-8"))
        digest.update(path.read_bytes())
    manifest["dataVersion"] = digest.hexdigest()
    _write_json(manifest_path, manifest)


def build_assets(source: Path, output_dir: Path, report_dir: Path) -> ImportResult:
    result = parse_workbook(source)
    grammar_path = output_dir / "grammar_points.json"
    question_path = output_dir / "tem4_questions.json"
    link_path = output_dir / "grammar_question_links.json"
    _write_json(grammar_path, result.grammar_points)
    _write_json(question_path, result.questions)
    _write_json(link_path, result.links)
    _write_json(report_dir / "grammar_tem4_audit_report.json", result.audit)
    _update_manifest(output_dir / "data_manifest.json", source, result, [grammar_path, question_path, link_path])
    return result


def main() -> int:
    root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description="Build normalized GrammarPoint/TEM4 assets from the source workbook.")
    parser.add_argument("--source", type=Path, default=root / "data-source/grammar-tem4/grammar_tem4_source.xlsx")
    parser.add_argument("--output", type=Path, default=root / "app/src/main/assets/database")
    parser.add_argument("--report-output", type=Path, default=root / "build/reports/data")
    args = parser.parse_args()
    result = build_assets(args.source, args.output, args.report_output)
    print(json.dumps(result.audit, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
